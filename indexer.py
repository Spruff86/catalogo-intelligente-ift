import os
import openpyxl
from pypdf import PdfReader
from pathlib import Path
from tqdm import tqdm
import database
import config

def build_file_map(directory: Path, valid_extensions: set):
    """Build a mapping of normalized SKU base to actual file path."""
    file_map = {}
    if not directory.exists():
        print(f"Warning: Directory not found: {directory}")
        return file_map
        
    for f in os.listdir(directory):
        if f.startswith('.'):
            continue
        ext = os.path.splitext(f)[1].lower()
        if ext in valid_extensions:
            full_path = directory / f
            base_name = os.path.splitext(f)[0].strip()
            
            # Map exact filename stem
            file_map[base_name.upper()] = str(full_path)
            
            # Also clean variations like _t, _s, _eng, etc.
            clean_name = base_name.replace("_t", "").replace("_s", "").replace("_eng", "").replace("__t", "").replace("_t_1", "").replace("_t ENG", "").strip().upper()
            if clean_name not in file_map:
                file_map[clean_name] = str(full_path)
                
    return file_map

def extract_pdf_text(pdf_path: str) -> str:
    """Extract text from a PDF file using pypdf."""
    try:
        reader = PdfReader(pdf_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts).strip()
    except Exception as e:
        return ""

def get_gdrive_id(file_path: str) -> str:
    """Extract Google Drive item ID from macOS extended attributes if available."""
    if not file_path:
        return ""
    try:
        import subprocess
        out = subprocess.check_output(
            ['xattr', '-p', 'com.google.drivefs.item-id#S', file_path],
            text=True,
            stderr=subprocess.DEVNULL
        ).strip()
        return out
    except Exception:
        return ""

def normalize_category(cat: str) -> str:
    """Cleans official category string from Excel and capitalizes first character for consistent UI presentation."""
    if not cat:
        return "Altro"
    c = str(cat).strip()
    return (c[0].upper() + c[1:]) if c else "Altro"

def load_historical_prices() -> dict:
    """Extracts known non-zero prices from previously indexed Excel files in the base directory."""
    import glob
    historical = {}
    excel_files = sorted(glob.glob(str(config.BASE_DIR / "Dati per catalogo*.xlsx")), reverse=True)
    for fn in excel_files:
        try:
            wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            h_idx = -1
            for i, r in enumerate(rows):
                if any(c in ['Cod_art', 'Item_Code'] for c in [str(x) for x in r if x]):
                    h_idx = i
                    break
            if h_idx == -1:
                continue
            header = [str(x).strip() if x is not None else "" for x in rows[h_idx]]
            s_idx = header.index('Soglia') if 'Soglia' in header else -1
            c_idx = header.index('Contratto') if 'Contratto' in header else -1
            sku_idx = header.index('Cod_art') if 'Cod_art' in header else (header.index('Item_Code') if 'Item_Code' in header else 2)
            if s_idx == -1 and c_idx == -1:
                continue
            for r in rows[h_idx + 1:]:
                sku = str(r[sku_idx] if sku_idx < len(r) else '').strip().upper()
                if not sku:
                    continue
                s_val = r[s_idx] if s_idx != -1 and s_idx < len(r) else 0
                c_val = r[c_idx] if c_idx != -1 and c_idx < len(r) else 0
                try:
                    sf = float(str(s_val or 0).replace(',', '.').strip())
                except (ValueError, TypeError):
                    sf = 0.0
                try:
                    cf = float(str(c_val or 0).replace(',', '.').strip())
                except (ValueError, TypeError):
                    cf = 0.0
                if (sf > 0 or cf > 0) and sku not in historical:
                    historical[sku] = (sf, cf)
        except Exception as e:
            print(f"Notice reading historical file {fn}: {e}")
    return historical

def run_indexing():
    print("Initializing database...")
    database.init_db()
    
    print("Scanning photo and technical sheet directories...")
    foto_map = build_file_map(config.FOTO_DIR, {'.jpg', '.jpeg', '.png'})
    schede_map = build_file_map(config.SCHEDE_DIR, {'.pdf'})
    
    print(f"Found {len(foto_map)} photo mappings and {len(schede_map)} technical sheet mappings.")
    
    print(f"Reading Excel catalog from {config.DATA_EXCEL_PATH}...")
    wb = openpyxl.load_workbook(config.DATA_EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        print("Excel is empty!")
        return
        
    # Find the header row (the first row containing 'Cod_art' or 'Item_Code')
    header_idx = -1
    for i, row in enumerate(all_rows):
        row_str = [str(c).strip() if c is not None else "" for c in row]
        if "Cod_art" in row_str or "Item_Code" in row_str or "Descrizione_articolo" in row_str:
            header_idx = i
            break
            
    if header_idx == -1:
        print("Could not find header row in Excel!")
        return
        
    header = [str(col).strip() if col is not None else "" for col in all_rows[header_idx]]
    data_rows = all_rows[header_idx + 1:]
    print(f"Found header at row {header_idx + 1} with {len(data_rows)} data rows. Columns: {header[:10]}...")
    
    # Map header names to indices
    col_map = {col: i for i, col in enumerate(header)}
    
    def get_val(row, name, default=""):
        idx = col_map.get(name)
        if idx is not None and idx < len(row) and row[idx] is not None:
            v = str(row[idx]).strip()
            if v == "zzz-zzz-zzz-zzz-zzz-zzz":
                return ""
            return v
        return default

    conn = database.get_connection()
    cursor = conn.cursor()
    
    # Load historical prices to guard against empty or zero price exports from SAP
    historical_prices = load_historical_prices()
    print(f"Loaded {len(historical_prices)} reference prices for price integrity guard.")
    
    # Cache existing media mappings, extracted PDF texts, and prices before clearing
    existing_media = {}
    try:
        cursor.execute("SELECT cod_art, gdrive_foto_id, gdrive_scheda_id, scheda_text, prezzo_indicativo, prezzo_contratto FROM products")
        for r in cursor.fetchall():
            existing_media[r[0]] = {
                "gdrive_foto_id": r[1] or "",
                "gdrive_scheda_id": r[2] or "",
                "scheda_text": r[3] or "",
                "prezzo_indicativo": r[4] or "0",
                "prezzo_contratto": r[5] or "0"
            }
    except Exception:
        pass

    # Clear existing data
    cursor.execute("DELETE FROM products")
    cursor.execute("DELETE FROM products_fts")
    
    products_to_insert = []
    
    print("Processing products and extracting PDF texts...")
    for row in tqdm(data_rows):
        cod_art = get_val(row, "Cod_art") or get_val(row, "Item_Code")
        if not cod_art:
            continue
            
        cod_art_clean = cod_art.strip().upper()
        
        descrizione_articolo = get_val(row, "Descrizione_articolo") or get_val(row, "Item_description")
        gruppo_art_raw = get_val(row, "Gruppo_art") or get_val(row, "Item_Group")
        gruppo_art = normalize_category(gruppo_art_raw)
        um = get_val(row, "Um") or get_val(row, "M.U.")
        tipo_peso = get_val(row, "Tipo_peso") or get_val(row, "Weight_note")
        unita_imballi = get_val(row, "Unità_e_imballi") or get_val(row, "Units_and_Packaging")
        ambiente_cons = get_val(row, "Ambiente_cons") or get_val(row, "Storage_Env")
        shelf_life_it = get_val(row, "Shelf_life_IT") or get_val(row, "Shelf_life_LNG")
        tempo_cons_gg = get_val(row, "TempoCons_gg") or get_val(row, "DeliveryTime_days")
        minimo_ordine = get_val(row, "Minimo_ordine") or get_val(row, "Minimum_order")
        descrizione_commerciale = get_val(row, "Descrizione Commerciale") or get_val(row, "Commercial_Description")
        
        # Official selling price: Prezzo Soglia
        prezzo_indicativo = get_val(row, "Soglia") or get_val(row, "Obiettivo") or get_val(row, "Contratto") or get_val(row, "Prz_BP")
        # Internal reference price: Prezzo Contratto
        prezzo_contratto = get_val(row, "Contratto") or get_val(row, "Prz_BP") or get_val(row, "Soglia")

        # Price Integrity Guard: if incoming price is zero or missing, fallback to valid historical price
        try:
            p_val = float(str(prezzo_indicativo or "0").replace(',', '.').strip())
        except (ValueError, TypeError):
            p_val = 0.0

        if p_val <= 0.0:
            cached_info = existing_media.get(cod_art, {})
            c_p = cached_info.get("prezzo_indicativo")
            if c_p and c_p not in ("0", "0.0", "", "None"):
                prezzo_indicativo = c_p
                prezzo_contratto = cached_info.get("prezzo_contratto") or c_p
            elif cod_art_clean in historical_prices:
                h_soglia, h_contratto = historical_prices[cod_art_clean]
                if h_soglia > 0:
                    prezzo_indicativo = str(h_soglia)
                if h_contratto > 0:
                    prezzo_contratto = str(h_contratto)

        art_attivo = get_val(row, "Art_Attivo", "Y")
        
        # Match Photo
        foto_path = foto_map.get(cod_art_clean, "")
        if not foto_path:
            # try variants
            base_sku = cod_art_clean.replace("PZ", "").replace("FRO", "").replace("CRT", "").replace("QU", "").strip()
            foto_path = foto_map.get(base_sku, "")
            
        foto_filename = os.path.basename(foto_path) if foto_path else ""
        
        # Match Scheda Tecnica
        scheda_path = schede_map.get(cod_art_clean, "")
        if not scheda_path:
            base_sku = cod_art_clean.replace("PZ", "").replace("FRO", "").replace("CRT", "").replace("QU", "").strip()
            scheda_path = schede_map.get(base_sku, "")
            
        scheda_filename = os.path.basename(scheda_path) if scheda_path else ""
        
        # Reuse cached PDF text and gdrive IDs if already available (fast indexing)
        cached = existing_media.get(cod_art, {})
        scheda_text = cached.get("scheda_text", "")
        if not scheda_text and scheda_path and os.path.exists(scheda_path):
            scheda_text = extract_pdf_text(scheda_path)
            
        gdrive_foto_id = cached.get("gdrive_foto_id") or get_gdrive_id(foto_path)
        gdrive_scheda_id = cached.get("gdrive_scheda_id") or get_gdrive_id(scheda_path)

        products_to_insert.append((
            cod_art,
            descrizione_articolo,
            gruppo_art,
            um,
            tipo_peso,
            unita_imballi,
            ambiente_cons,
            shelf_life_it,
            tempo_cons_gg,
            minimo_ordine,
            descrizione_commerciale,
            prezzo_indicativo,
            prezzo_contratto,
            art_attivo,
            foto_filename,
            foto_path,
            scheda_filename,
            scheda_path,
            scheda_text,
            gdrive_foto_id,
            gdrive_scheda_id
        ))
        
    print(f"Inserting {len(products_to_insert)} products into database...")
    cursor.executemany("""
    INSERT INTO products (
        cod_art,
        descrizione_articolo,
        gruppo_art,
        um,
        tipo_peso,
        unita_imballi,
        ambiente_cons,
        shelf_life_it,
        tempo_cons_gg,
        minimo_ordine,
        descrizione_commerciale,
        prezzo_indicativo,
        prezzo_contratto,
        art_attivo,
        foto_filename,
        foto_path,
        scheda_filename,
        scheda_path,
        scheda_text,
        gdrive_foto_id,
        gdrive_scheda_id
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, products_to_insert)
    
    # Populate FTS
    print("Building full-text search index...")
    cursor.execute("""
    INSERT INTO products_fts(rowid, cod_art, descrizione_articolo, gruppo_art, descrizione_commerciale, scheda_text)
    SELECT rowid, cod_art, descrizione_articolo, gruppo_art, descrizione_commerciale, scheda_text FROM products
    """)
    
    conn.commit()
    conn.close()
    print("Indexing complete! Database ready.")

if __name__ == "__main__":
    run_indexing()
